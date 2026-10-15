"""Пересобирает utm-data.js из таблицы «Генератор UTM VASTU HOME» (.xlsx).

Запуск: python3 tools/update_utm_data.py путь/к/таблице.xlsx
Нужен пакет openpyxl (pip install openpyxl).
"""
import json, sys
import openpyxl

def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ref = wb["справочник"]

    def pairs(col_name, col_val):
        out = []
        for row in ref.iter_rows(min_row=3):
            name, val = row[col_name].value, row[col_val].value
            if name is None or val is None:
                continue
            out.append([str(name).strip(), str(val).strip()])
        return out

    landings = []
    for row in wb["ID лендингов"].iter_rows(min_row=3):
        name, url, _, _, lid = [c.value for c in row[:5]]
        if name and url:
            landings.append([str(name).strip(), str(url).strip(), str(lid or "").strip()])

    data = {
        "sources": pairs(0, 1),    # Источник трафика → utm_source
        "channels": pairs(2, 3),   # Канал привлечения → utm_medium
        "types": pairs(4, 5),      # Тип контента → utm_campaign
        "landings": landings,      # Страница, ссылка, ID лендинга → utm_term
    }
    with open("utm-data.js", "w", encoding="utf-8") as f:
        f.write("/* Справочники генератора UTM VASTU HOME. Собрано tools/update_utm_data.py из таблицы. */\n")
        f.write("window.UTM_DATA = " + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n")
    print(f"utm-data.js: {len(data['sources'])} источников, {len(data['channels'])} каналов, "
          f"{len(data['types'])} типов контента, {len(landings)} лендингов")

if __name__ == "__main__":
    main(sys.argv[1])
